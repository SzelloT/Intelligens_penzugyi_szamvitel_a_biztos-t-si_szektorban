"""Streamlit alapú, webes felületű számlafeldolgozó alkalmazás."""

import json
import logging
import tempfile
import time
from datetime import datetime
from pathlib import Path
from typing import Optional
import warnings

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

warnings.filterwarnings("ignore")
logging.getLogger("google").setLevel(logging.ERROR)

# ==============================================================================
# KONFIGURÁCIÓ
# ==============================================================================
MODELS_TO_USE = [
    "gemini-3.8-flash",
    "gemini-flash-latest"
]

REQUEST_DELAY_SECONDS = 15

# Streamlit oldal beállítása
st.set_page_config(page_title="Számlafeldolgozó", layout="centered")


# ==============================================================================
# ADATMODELL
# ==============================================================================
class InvoiceData(BaseModel):
    barcode: Optional[str] = Field(default=None, description="Vonalkód értéke vagy száma")
    supplier_name: Optional[str] = Field(default=None, description="Szállító neve")
    supplier_tax_number: Optional[str] = Field(default=None, description="Szállító adószáma")
    supplier_address: Optional[str] = Field(default=None, description="Szállító címe")
    supplier_bank_account: Optional[str] = Field(default=None, description="Szállító bankszámlaszáma")
    invoice_number: Optional[str] = Field(default=None, description="Számlaszám")
    issue_date: Optional[str] = Field(default=None, description="Számla kelte (ÉÉÉÉ.HH.NN)")
    fulfillment_date: Optional[str] = Field(default=None, description="Teljesítés dátuma (ÉÉÉÉ.HH.NN)")
    payment_due_date: Optional[str] = Field(default=None, description="Fizetési határidő (ÉÉÉÉ.HH.NN)")
    currency: Optional[str] = Field(default=None, description="Pénznem (pl. HUF, EUR)")
    net_amount: Optional[str] = Field(default=None, description="Nettó összeg")
    vat_amount: Optional[str] = Field(default=None, description="Áfa összege")
    gross_amount: Optional[str] = Field(default=None, description="Bruttó összeg")


# ==============================================================================
# ADATKINYERÉSI LOGIKA
# ==============================================================================
def extract_invoice_data(client: genai.Client, file_path: Path, max_retries: int = 3) -> Optional[InvoiceData]:
    uploaded_file = None
    try:
        uploaded_file = client.files.upload(file=str(file_path))
        prompt = (
            "Olvasd ki a csatolt számláról a kért 13 adatmezőt. "
            "A választ pontosan a megadott séma szerint JSON formátumban add vissza. "
            "Amennyiben egy adat hiányzik a számlán (pl. nincs vonalkód), értéke legyen null."
        )

        for model_name in MODELS_TO_USE:
            for attempt in range(1, max_retries + 1):
                try:
                    chat = client.chats.create(
                        model=model_name,
                        config=types.GenerateContentConfig(
                            response_mime_type="application/json",
                            response_schema=InvoiceData,
                            temperature=0.1,
                        ),
                    )
                    response = chat.send_message(message=[uploaded_file, prompt])

                    if hasattr(response, "parsed") and response.parsed is not None:
                        return response.parsed

                    if response.text:
                        clean_text = response.text.strip()
                        if clean_text.startswith("```json"):
                            clean_text = clean_text[7:]
                        if clean_text.startswith("```"):
                            clean_text = clean_text[3:]
                        if clean_text.endswith("```"):
                            clean_text = clean_text[:-3]

                        parsed_json = json.loads(clean_text.strip())
                        return InvoiceData.model_validate(parsed_json)
                    return None

                except Exception as exc:
                    err_msg = str(exc)
                    if "503" in err_msg or "UNAVAILABLE" in err_msg:
                        time.sleep(10)
                        if attempt == max_retries:
                            break
                    elif "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
                        time.sleep(25)
                        if attempt == max_retries:
                            break
                    else:
                        break
        return None
    finally:
        if uploaded_file:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass


def sanitize_invoice_data(data: Optional[InvoiceData], file_name: str) -> InvoiceData:
    if data is None:
        data = InvoiceData()
    invalid_literals = {"none", "null", "n/a", "nincs", "-", ""}
    current_barcode = str(data.barcode).strip() if data.barcode else ""
    if not current_barcode or current_barcode.lower() in invalid_literals:
        data.barcode = Path(file_name).stem
    return data


# ==============================================================================
# EXCEL GENERÁLÁS (SABLON NÉLKÜL)
# ==============================================================================
FIELD_DEFINITIONS = [
    ("barcode", "Vonalkód:"),
    ("supplier_name", "Szálító neve:"),
    ("supplier_tax_number", "Szállító adószáma:"),
    ("supplier_address", "Szállító címe:"),
    ("supplier_bank_account", "Szállító bankszámlaszáma:"),
    ("invoice_number", "Számla számlaszáma:"),
    ("issue_date", "Számla kelte:"),
    ("fulfillment_date", "Számla teljesítési dátuma:"),
    ("payment_due_date", "Számla fizetési határidő:"),
    ("currency", "Számla pénzneme:"),
    ("net_amount", "Számla nettó összege:"),
    ("vat_amount", "Számla áfa:"),
    ("gross_amount", "Számla bruttó összege:"),
]


def create_excel_with_results(output_path: Path, results: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Szamlak"
    ws.views.sheetView[0].showGridLines = True

    # Stílusok
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    label_font = Font(name="Calibri", size=11, bold=True)

    # Hiba formázás: piros kitöltés, fehér félkövér betű
    error_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    error_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    header_border = Border(
        left=Side(style="thin", color="FFFFFF"),
        right=Side(style="thin", color="FFFFFF"),
        top=Side(style="thin", color="FFFFFF"),
        bottom=Side(style="thin", color="FFFFFF"),
    )

    headers = {
        1: ("Nr.", Alignment(horizontal="center", vertical="center")),
        2: ("Időpont", Alignment(horizontal="center", vertical="center")),
        3: ("Adat típus", Alignment(horizontal="left", vertical="center")),
        4: ("Adatok", Alignment(horizontal="left", vertical="center")),
        5: ("Megjegyzés", Alignment(horizontal="center", vertical="center")),
    }

    # Fejléc kiírása
    ws.row_dimensions[1].height = 26
    for col_idx, (text, align) in headers.items():
        cell = ws.cell(row=1, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align
        cell.border = header_border

    invalid_literals = {"", "none", "null", "n/a", "nincs", "-"}

    # Sorok kiírása
    for inv_idx, (proc_time, data, file_name, is_success) in enumerate(results, start=1):
        start_row = 2 + (inv_idx - 1) * len(FIELD_DEFINITIONS)
        time_str = proc_time.strftime("%Y.%m.%d %H:%M:%S")

        for offset, (field_key, field_label) in enumerate(FIELD_DEFINITIONS):
            current_row = start_row + offset
            ws.row_dimensions[current_row].height = 20

            raw_val = getattr(data, field_key, None)
            val_str = str(raw_val).strip() if raw_val is not None else ""
            if val_str.lower() in invalid_literals:
                val_str = ""

            # Hiba ellenőrzése: üres érték vagy sikertelen hívás
            is_field_error = (val_str == "") or (not is_success)

            if is_field_error:
                comment_text = f"Hiba! (Forrás: {file_name})" if offset == 0 else "Hiba!"
                comment_font = error_font
                comment_fill = error_fill
                comment_align = Alignment(horizontal="center", vertical="center")
            else:
                comment_text = f"Forrás: {file_name}" if offset == 0 else ""
                comment_font = data_font
                comment_fill = PatternFill(fill_type=None)
                comment_align = Alignment(horizontal="left", vertical="center")

            row_values = {
                1: (inv_idx, Alignment(horizontal="center", vertical="center"), data_font, None),
                2: (time_str, Alignment(horizontal="center", vertical="center"), data_font, None),
                3: (field_label, Alignment(horizontal="left", vertical="center"), label_font, None),
                4: (val_str, Alignment(horizontal="left", vertical="center"), data_font, None),
                5: (comment_text, comment_align, comment_font, comment_fill),
            }

            for col_idx, (val, align, font_style, fill_style) in row_values.items():
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_style
                cell.alignment = align
                cell.border = thin_border
                if fill_style:
                    cell.fill = fill_style

    # Oszlopszélességek beállítása
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 35

    wb.save(output_path)


# ==============================================================================
# STREAMLIT FELÜLET
# ==============================================================================
st.title("🧾 Számlafeldolgozó AI")

api_key = st.text_input("Gemini API Kulcs", type="password", placeholder="Ide illeszd be az API kulcsod...")

uploaded_files = st.file_uploader(
    "Húzd ide a feldolgozandó számlákat",
    accept_multiple_files=True,
    type=["pdf", "jpg", "jpeg", "png"],
)

if st.button("Feldolgozás indítása", type="primary"):
    if not api_key:
        st.warning("Kérlek, add meg az API kulcsot!")
        st.stop()
    if not uploaded_files:
        st.warning("Kérlek, tölts fel legalább egy számlát!")
        st.stop()

    client = genai.Client(api_key=api_key)
    results = []
    total_files = len(uploaded_files)

    progress_bar = st.progress(0)
    status_text = st.empty()

    for idx, uploaded_file in enumerate(uploaded_files, start=1):
        status_text.text(f"Feldolgozás alatt ({idx}/{total_files}): {uploaded_file.name}...")

        with tempfile.NamedTemporaryFile(delete=False, suffix=Path(uploaded_file.name).suffix) as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = Path(tmp_file.name)

        proc_time = datetime.now()
        raw_data = extract_invoice_data(client, tmp_file_path)
        is_success = raw_data is not None

        processed_data = sanitize_invoice_data(raw_data, uploaded_file.name)
        results.append((proc_time, processed_data, uploaded_file.name, is_success))

        tmp_file_path.unlink(missing_ok=True)

        progress_bar.progress(idx / total_files)

        if idx < total_files:
            status_text.text(f"Várakozás a kvóta miatt ({REQUEST_DELAY_SECONDS} mp)...")
            time.sleep(REQUEST_DELAY_SECONDS)

    status_text.text("Excel fájl generálása...")

    run_timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_filename = f"Kimeneti_szamlak_{run_timestamp}.xlsx"

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_excel:
        temp_excel_path = Path(tmp_excel.name)

    create_excel_with_results(temp_excel_path, results)

    with open(temp_excel_path, "rb") as f:
        excel_data = f.read()

    st.success("Feldolgozás sikeresen befejeződött!")
    st.download_button(
        label="📥 Kész Excel fájl letöltése",
        data=excel_data,
        file_name=output_filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    temp_excel_path.unlink(missing_ok=True)