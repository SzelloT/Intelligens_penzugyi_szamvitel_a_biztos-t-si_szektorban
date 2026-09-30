"""Streamlit alapú, webes felületű számlafeldolgozó alkalmazás."""

import json
import logging
import os
import tempfile
import time
from datetime import datetime
from pathlib import Path
from typing import Optional
import warnings

import openpyxl
import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel, Field

warnings.filterwarnings("ignore")
logging.getLogger("google").setLevel(logging.ERROR)

# ==============================================================================
# KONFIGURÁCIÓ
# ==============================================================================
MODELS_TO_USE = [
    "gemini-3.8-flash",
    "gemini-flash-latest"
]

TEMPLATE_EXCEL = Path("szamla_sablon.xlsx")
REQUEST_DELAY_SECONDS = 15

# Streamlit oldal beállítása
st.set_page_config(page_title="Számlafeldolgozó", layout="centered")


# ==============================================================================
# ADATMODELL
# ==============================================================================
class InvoiceData(BaseModel):
    barcode: Optional[str] = Field(default=None, description="Vonalkód értéke vagy száma")
    supplier_name: Optional[str] = Field(default=None, description="Szállító neve")
    supplier_tax_number: Optional[str] = Field(default=None, description="Szállító adószáma")
    supplier_address: Optional[str] = Field(default=None, description="Szállító címe")
    supplier_bank_account: Optional[str] = Field(default=None, description="Szállító bankszámlaszáma")
    invoice_number: Optional[str] = Field(default=None, description="Számlaszám")
    issue_date: Optional[str] = Field(default=None, description="Számla kelte (ÉÉÉÉ.HH.NN)")
    fulfillment_date: Optional[str] = Field(default=None, description="Teljesítés dátuma (ÉÉÉÉ.HH.NN)")
    payment_due_date: Optional[str] = Field(default=None, description="Fizetési határidő (ÉÉÉÉ.HH.NN)")
    currency: Optional[str] = Field(default=None, description="Pénznem (pl. HUF, EUR)")
    net_amount: Optional[str] = Field(default=None, description="Nettó összeg")
    vat_amount: Optional[str] = Field(default=None, description="Áfa összege")
    gross_amount: Optional[str] = Field(default=None, description="Bruttó összeg")


# ==============================================================================
# ADATKINYERÉSI LOGIKA
# ==============================================================================
def extract_invoice_data(client: genai.Client, file_path: Path, max_retries: int = 3) -> Optional[InvoiceData]:
    uploaded_file = None
    try:
        uploaded_file = client.files.upload(file=str(file_path))
        prompt = (
            "Olvasd ki a csatolt számláról a kért 13 adatmezőt. "
            "A választ pontosan a megadott séma szerint JSON formátumban add vissza. "
            "Amennyiben egy adat hiányzik a számlán (pl. nincs vonalkód), értéke legyen null."
        )

        for model_name in MODELS_TO_USE:
            for attempt in range(1, max_retries + 1):
                try:
                    chat = client.chats.create(
                        model=model_name,
                        config=types.GenerateContentConfig(
                            response_mime_type="application/json",
                            response_schema=InvoiceData,
                            temperature=0.1,
                        ),
                    )
                    response = chat.send_message(message=[uploaded_file, prompt])

                    if hasattr(response, "parsed") and response.parsed is not None:
                        return response.parsed

                    if response.text:
                        clean_text = response.text.strip()
                        if clean_text.startswith("```json"):
                            clean_text = clean_text[7:]
                        if clean_text.startswith("```"):
                            clean_text = clean_text[3:]
                        if clean_text.endswith("```"):
                            clean_text = clean_text[:-3]

                        parsed_json = json.loads(clean_text.strip())
                        return InvoiceData.model_validate(parsed_json)
                    return None

                except Exception as exc:
                    err_msg = str(exc)
                    if "503" in err_msg or "UNAVAILABLE" in err_msg:
                        time.sleep(10)
                        if attempt == max_retries:
                            break
                    elif "429" in err_msg or "RESOURCE_EXHAUSTED" in err_msg:
                        time.sleep(25)
                        if attempt == max_retries:
                            break
                    else:
                        break
        return None
    finally:
        if uploaded_file:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass


def sanitize_invoice_data(data: Optional[InvoiceData], file_name: str) -> InvoiceData:
    if data is None:
        data = InvoiceData()
    invalid_literals = {"none", "null", "n/a", "nincs", "-", ""}
    current_barcode = str(data.barcode).strip() if data.barcode else ""
    if not current_barcode or current_barcode.lower() in invalid_literals:
        data.barcode = Path(file_name).stem
    return data


def write_to_excel(template_path: Path, output_path: Path, results: list):
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Szamlak"]

    fields = [
        "barcode", "supplier_name", "supplier_tax_number", "supplier_address",
        "supplier_bank_account", "invoice_number", "issue_date", "fulfillment_date",
        "payment_due_date", "currency", "net_amount", "vat_amount", "gross_amount",
    ]

    for idx, (proc_time, data, file_name, is_success) in enumerate(results, start=1):
        start_row = 2 + (idx - 1) * 13
        time_str = proc_time.strftime("%Y.%m.%d %H:%M:%S")

        for offset, field_name in enumerate(fields):
            current_row = start_row + offset
            val = getattr(data, field_name, None)

            ws.cell(row=current_row, column=2, value=time_str)

            clean_val = str(val).strip() if val is not None else ""
            if clean_val.lower() in {"none", "null"}:
                clean_val = ""
            ws.cell(row=current_row, column=4, value=clean_val)

            if offset == 0:
                status_suffix = "" if is_success else " (HIBA)"
                ws.cell(row=current_row, column=5, value=f"Forrás: {file_name}{status_suffix}")

    wb.save(output_path)


# ==============================================================================
# STREAMLIT FELÜLET
# ==============================================================================
st.title("🧾 Számlafeldolgozó AI")

if not TEMPLATE_EXCEL.exists():
    st.error(f"HIBA: A sablonfájl ({TEMPLATE_EXCEL.name}) nem található a mappában!")
    st.stop()

api_key = st.text_input("Gemini API Kulcs", type="password", placeholder="Ide illeszd be az API kulcsod...")

uploaded_files = st.file_uploader(
    "Húzd ide a feldolgozandó számlákat",
    accept_multiple_files=True,
    type=["pdf", "jpg", "jpeg", "png"]
)

if st.button("Feldolgozás indítása", type="primary"):
    if not api_key:
        st.warning("Kérlek, add meg az API kulcsot!")
        st.stop()
    if not uploaded_files:
        st.warning("Kérlek, tölts fel legalább egy számlát!")
        st.stop()

    client = genai.Client(api_key=api_key)
    results = []
    total_files = len(uploaded_files)

    progress_bar = st.progress(0)
    status_text = st.empty()

    for idx, uploaded_file in enumerate(uploaded_files, start=1):
        status_text.text(f"Feldolgozás alatt ({idx}/{total_files}): {uploaded_file.name}...")

        # Fájl ideiglenes mentése, hogy az SDK fel tudja tölteni
        with tempfile.NamedTemporaryFile(delete=False, suffix=Path(uploaded_file.name).suffix) as tmp_file:
            tmp_file.write(uploaded_file.read())
            tmp_file_path = Path(tmp_file.name)

        proc_time = datetime.now()
        raw_data = extract_invoice_data(client, tmp_file_path)
        is_success = raw_data is not None

        processed_data = sanitize_invoice_data(raw_data, uploaded_file.name)
        results.append((proc_time, processed_data, uploaded_file.name, is_success))

        # Ideiglenes fájl törlése
        tmp_file_path.unlink(missing_ok=True)

        progress_bar.progress(idx / total_files)

        if idx < total_files:
            status_text.text(f"Várakozás a kvóta miatt ({REQUEST_DELAY_SECONDS} mp)...")
            time.sleep(REQUEST_DELAY_SECONDS)

    status_text.text("Excel fájl generálása...")

    # Eredmény mentése egy ideiglenes Excel fájlba
    run_timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    output_filename = f"Kimeneti_szamlak_{run_timestamp}.xlsx"

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_excel:
        temp_excel_path = Path(tmp_excel.name)

    write_to_excel(TEMPLATE_EXCEL, temp_excel_path, results)

    # Letöltés gomb megjelenítése
    with open(temp_excel_path, "rb") as f:
        excel_data = f.read()

    st.success("Feldolgozás sikeresen befejeződött!")
    st.download_button(
        label="📥 Kész Excel fájl letöltése",
        data=excel_data,
        file_name=output_filename,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    # Ideiglenes Excel fájl törlése
    temp_excel_path.unlink(missing_ok=True)